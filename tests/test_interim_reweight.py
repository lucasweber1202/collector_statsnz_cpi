"""Regression coverage for official interim baskets, regime boundaries and stored reconciliation."""

from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from scripts.extract import SourceData, SourceLayoutError
from scripts.reconcile import reconcile
from scripts.target_weights import derive_weights
from scripts.time_series import Observation
from scripts.weight_sources import BaseWeight, parse_interim_weights

ROOT = "STATSNZ_CPI_CPIQ_SE9A"
PARENT = "STATSNZ_CPI_CPIQ_SE9011"
A, B = "STATSNZ_CPI_CPIQ_SE901101", "STATSNZ_CPI_CPIQ_SE901102"


def test_corrected_interim_workbook_preserves_every_percentage() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Table 1"
    for col, year in [(11, 2021), (13, 2022), (15, 2023)]:
        sheet.cell(8, col, f"June {year}")
    catalog = {ROOT: {"name": "All groups"}}
    for i in range(1, 12):
        sid = f"STATSNZ_CPI_CPIQ_SE9{i:02d}"
        catalog[sid] = {"name": f"Group {i}"}
        sheet.cell(10 + i, 1, f"Group {i} group")
        for col in (11, 13, 15):
            sheet.cell(10 + i, col, 9.09 if i < 11 else 9.1)
    buffer = io.BytesIO()
    book.save(buffer)
    weights = parse_interim_weights(buffer.getvalue(), catalog)
    assert len(weights) == 36
    assert {w.base_period for w in weights} == {date(y, 6, 30) for y in (2021, 2022, 2023)}
    assert next(w.percent for w in weights if w.series_id.endswith("SE901")) == 9.09
    sheet.cell(8, 15, "Unexpected base")
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(SourceLayoutError, match="headers"):
        parse_interim_weights(buffer.getvalue(), catalog)


def _base(sid: str, parent: str | None, base: date, weight: float) -> BaseWeight:
    return BaseWeight(sid, parent, base, weight / 100, None, sid, weight)


def test_opening_shares_switch_strictly_after_official_base() -> None:
    old, new, before, after = (
        date(2020, 6, 30),
        date(2021, 6, 30),
        date(2021, 3, 31),
        date(2021, 9, 30),
    )
    originals = [
        w
        for base, wa in [(old, 50), (new, 20)]
        for w in (
            _base(PARENT, "STATSNZ_CPI_CPIQ_SE901", base, 100),
            _base(A, PARENT, base, wa),
            _base(B, PARENT, base, 100 - wa),
        )
    ]
    levels = {old: (100, 100), before: (200, 100), new: (200, 100), after: (400, 100)}
    points = [
        Observation(sid, t, value, "fixture")
        for t, pair in levels.items()
        for sid, value in [(A, pair[0]), (B, pair[1]), (ROOT, 100)]
    ]
    output = derive_weights(
        SourceData(points, {sid: {} for sid in (A, B, PARENT, ROOT)}), originals
    )
    shares = {(w.series_id, w.reference_date): w.weight for w in output}
    assert shares[A, new] == pytest.approx(2 / 3)  # old basket, opening March levels
    assert shares[A, after] == pytest.approx(
        0.2
    )  # official June basket, not closing September levels


def test_missing_positive_component_never_renormalizes_remaining_sibling() -> None:
    base, current = date(2020, 6, 30), date(2020, 9, 30)
    points = [Observation(sid, t, 100, "fixture") for sid in (A, ROOT) for t in (base, current)]
    originals = [
        _base(PARENT, "STATSNZ_CPI_CPIQ_SE901", base, 100),
        _base(A, PARENT, base, 50),
        _base(B, PARENT, base, 50),
    ]
    output = derive_weights(
        SourceData(points, {sid: {} for sid in (A, B, PARENT, ROOT)}), originals
    )
    assert not any(w.series_id in (A, B) for w in output)


def test_reconciliation_detects_wrong_weights_and_material_parent_error() -> None:
    base, current = date(2020, 6, 30), date(2020, 9, 30)
    levels = {(sid, base): 100.0 for sid in (A, B, PARENT)}
    levels.update({(A, current): 110.0, (B, current): 100.0, (PARENT, current): 105.0})
    shares = {(A, current): 0.5, (B, current): 0.5}
    originals = {(A, base): 50.0, (B, base): 50.0}
    parents = {A: PARENT, B: PARENT, PARENT: None}
    assert (
        reconcile(levels, shares, originals, parents)[-1]["status"] == "PUBLIC_PRECISION_COMPATIBLE"
    )
    levels[PARENT, current] = 120.0
    assert reconcile(levels, shares, originals, parents)[-1]["status"] == "UNEXPLAINED_RESIDUAL"
    with pytest.raises(ValueError, match="formula"):
        reconcile(levels, {(A, current): 0.6, (B, current): 0.4}, originals, parents)
