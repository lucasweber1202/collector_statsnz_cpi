"""Export stored target tables into an audit workbook (one sheet per table).

Exports latest vintages from the database, not a fresh external download, so
index levels and both weighting systems are exactly the records being audited.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
from sqlalchemy import text

from scripts.config import SCHEMA_NAME
from scripts.db import build_engine


def export_workbook(path: Path) -> None:
    """Write latest stored levels, effective shares, official baskets and hierarchy."""
    engine = build_engine()
    book = openpyxl.Workbook()
    active = book.active
    assert active is not None
    book.remove(active)
    try:
        with engine.connect() as conn:
            for table, measure in [
                ("time_series", "value"),
                ("weights", "weight"),
                ("original_weights", "weight"),
            ]:
                rows = (
                    conn.execute(
                        text(
                            f"SELECT * FROM (SELECT *, ROW_NUMBER() OVER (PARTITION BY series_id, reference_date ORDER BY vintage_date DESC,collected_at DESC) AS rn FROM {SCHEMA_NAME}.{table}) ranked WHERE rn=1 ORDER BY reference_date,series_id"
                        )
                    )
                    .mappings()
                    .all()
                )
                sheet = book.create_sheet(table)
                if table == "original_weights":
                    columns = [
                        "series_id",
                        "reference_date",
                        "vintage_date",
                        "weight",
                        "weight_base_year",
                        "collected_at",
                    ]
                    sheet.append(columns)
                    for row in rows:
                        sheet.append([row[c] for c in columns])
                else:
                    series = sorted({str(row["series_id"]) for row in rows})
                    values = {
                        (row["reference_date"], row["series_id"]): row[measure] for row in rows
                    }
                    sheet.append(["reference_date", *series])
                    for when in sorted({row["reference_date"] for row in rows}):
                        sheet.append([when, *[values.get((when, sid)) for sid in series]])
                sheet.freeze_panes = "B2"
            sheet = book.create_sheet("cpi_hierarchy")
            columns = ["series_id", "native_code", "parent_id", "level", "name", "collected_at"]
            sheet.append(columns)
            for row in conn.execute(
                text(f"SELECT * FROM {SCHEMA_NAME}.cpi_hierarchy ORDER BY level,series_id")
            ).mappings():
                sheet.append([row[c] for c in columns])
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(path)
    finally:
        engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser(description="Export stored CPI target validation tables")
    parser.add_argument("--output", type=Path, required=True)
    export_workbook(parser.parse_args().output)


if __name__ == "__main__":
    main()
